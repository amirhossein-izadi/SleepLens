import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_persian_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors matched to the HTML/PPTX design
    BG_DARK = RGBColor(11, 15, 23)        # #0b0f17
    CARD_BG = RGBColor(30, 41, 59)        # #1e293b
    CARD_BORDER = RGBColor(51, 65, 85)    # #334155
    TEXT_WHITE = RGBColor(248, 250, 252)  # #f8fafc
    TEXT_MUTED = RGBColor(203, 213, 225)  # #cbd5e1
    TEXT_SUBTLE = RGBColor(148, 163, 184) # #94a3b8
    TEAL_PRIMARY = RGBColor(20, 184, 166) # #14b8a6
    CYAN_ACCENT = RGBColor(6, 182, 212)   # #06b6d4
    EMERALD_GREEN = RGBColor(16, 185, 129)# #10b981
    AMBER_WARN = RGBColor(245, 158, 11)   # #f59e0b
    ROSE_ACCENT = RGBColor(244, 63, 94)   # #f43f5e
    WIN_HIGHLIGHT_BG = RGBColor(19, 78, 74) # #134e4a
    INNER_CARD_BG = RGBColor(15, 23, 42)  # #0f172a

    FONT_FA = "Tahoma"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category, subtitle="", timing=""):
        # Header text (right aligned for Persian)
        tb = slide.shapes.add_textbox(Inches(3.5), Inches(0.35), Inches(9.0), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.alignment = PP_ALIGN.RIGHT
        p_cat.text = category
        p_cat.font.name = FONT_FA
        p_cat.font.size = Pt(11.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL_PRIMARY
        
        p_title = tf.add_paragraph()
        p_title.alignment = PP_ALIGN.RIGHT
        p_title.text = title
        p_title.font.name = FONT_FA
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.space_before = Pt(3)

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.alignment = PP_ALIGN.RIGHT
            p_sub.text = subtitle
            p_sub.font.name = FONT_FA
            p_sub.font.size = Pt(12.5)
            p_sub.font.color.rgb = CYAN_ACCENT
            p_sub.space_before = Pt(3)

        # Timing Badge Pill on Top Left (RTL balance)
        if timing:
            pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.533), Inches(0.48))
            pill.fill.solid()
            pill.fill.fore_color.rgb = RGBColor(8, 47, 73)
            pill.line.color.rgb = RGBColor(14, 116, 144)
            pill.line.width = Pt(1)
            ptf = pill.text_frame
            ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0
            pp = ptf.paragraphs[0]
            pp.alignment = PP_ALIGN.CENTER
            pp.text = timing
            pp.font.name = FONT_FA
            pp.font.size = Pt(11)
            pp.font.bold = True
            pp.font.color.rgb = CYAN_ACCENT

    def add_card(slide, left, top, width, height, bg=CARD_BG, border=CARD_BORDER, border_width=1):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg
        card.line.color.rgb = border
        card.line.width = Pt(border_width)
        return card

    def add_footer(slide, right_text, left_text):
        tb = slide.shapes.add_textbox(Inches(4.0), Inches(6.92), Inches(8.533), Inches(0.35))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = right_text
        p.font.name = FONT_FA
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_SUBTLE

        tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(3.5), Inches(0.35))
        tf_l = tb_l.text_frame
        tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
        pl = tf_l.paragraphs[0]
        pl.alignment = PP_ALIGN.LEFT
        pl.text = left_text
        pl.font.name = FONT_FA
        pl.font.size = Pt(10.5)
        pl.font.color.rgb = TEXT_SUBTLE

    def add_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # =========================================================================
    # SLIDE 1: Title & Executive Summary (FA)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.733), Inches(2.2))
    ttf = t_box.text_frame
    ttf.word_wrap = True
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0

    p = ttf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "خلاصه اجرایی • ارائه نهایی پروژه هکاتون"
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = ttf.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "اسلیپ‌لنز: پلتفرم هوش مصنوعی دو‌مرحله‌ای برای دسته‌بندی استیج‌های خواب\nو سنجش شاخص پیوسته کیفیت خواب"
    p.font.name = FONT_FA
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(6)

    p = ttf.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "پیوند میان استیجینگ گسسته ۳۰ ثانیه‌ای با عمق پیوسته خواب و نشانگرهای ریزساختاری مغز"
    p.font.name = FONT_FA
    p.font.size = Pt(14)
    p.font.color.rgb = CYAN_ACCENT
    p.space_before = Pt(6)

    # 3 Cards RTL order: Right is Card 1, Middle is Card 2, Left is Card 3
    # Card 1 (Right): Amber border
    add_card(s1, Inches(8.88), Inches(2.95), Inches(3.64), Inches(3.8), border=AMBER_WARN, border_width=2)
    tb1 = s1.shapes.add_textbox(Inches(9.08), Inches(3.15), Inches(3.24), Inches(3.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0
    p = tf1.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "⚠️ گلوگاه بالینی در تست خواب"
    p.font.name = FONT_FA
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = AMBER_WARN

    bullets1 = [
        "• تحلیل دستی و چشمی سیگنال‌های پلی‌سومنوگرافی (PSG) نیازمند ۱.۵ تا ۲ ساعت زمان برای هر شب است.",
        "• عدم توافق بالای پزشکان در تشخیص مراحل گذار (دقت تشخیصی مرحله N1 کمتر از ۶۰٪ است).",
        "• سیستم سنتی ۵ مرحله‌ای (AASM) پیوستگی دینامیک عمق خواب و همگام‌سازی مغز را کاملا نادیده می‌گیرد."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = b
        p.font.name = FONT_FA
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(12)

    # Card 2 (Middle): Cyan border
    add_card(s1, Inches(4.84), Inches(2.95), Inches(3.64), Inches(3.8), border=CYAN_ACCENT, border_width=2)
    tb2 = s1.shapes.add_textbox(Inches(5.04), Inches(3.15), Inches(3.24), Inches(3.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "🧠 موتور هوش مصنوعی دو‌مرحله‌ای"
    p.font.name = FONT_FA
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    bullets2 = [
        "• مرحله اول (استیجینگ): انسمبل ۴ مدله با امتیاز F1 ماکرو 0.8425 (ضریب کاپای 0.831) روی ۲۳۷ هزار ایپاک.",
        "• مرحله دوم (عمق و کیفیت): ترنسفورمر SDI (منتشر شده در npj Digital Med) برای عمق پیوسته + ۹۶ پارامتر بالینی.",
        "• حداقل‌سازی سنسورها: چیدمان ۲ سنسوره (Fpz+EOG) توانسته ۹۹.۵٪ دقت سیستم کامل را حفظ کند!"
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = b
        p.font.name = FONT_FA
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(12)

    # Card 3 (Left): Emerald border
    add_card(s1, Inches(0.8), Inches(2.95), Inches(3.64), Inches(3.8), border=EMERALD_GREEN, border_width=2)
    tb3 = s1.shapes.add_textbox(Inches(1.0), Inches(3.15), Inches(3.24), Inches(3.4))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0
    p = tf3.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "💻 پلتفرم عملیاتی و بالینی"
    p.font.name = FONT_FA
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN

    bullets3 = [
        "• وب‌اپلیکیشن فول‌استک بالینی با فریم‌ورک‌های مدرن Next.js 14 و Django REST Framework.",
        "• هیپنوگرام تعاملی، منحنی عمق پیوسته خواب (SDI)، و کاوشگر لحظه‌ای سیگنال‌های خام مغزی.",
        "• تولید خودکار گزارش متنی پزشکی با LLM + چت‌بات دستیار تشخیصی مجهز به OpenCode."
    ]
    for b in bullets3:
        p = tf3.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = b
        p.font.name = FONT_FA
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(12)

    add_footer(s1, "پلتفرم اسلیپ‌لنز • نمونه اولیه سامانه هوش مصنوعی بالینی و تحقیقاتی خواب", "اسلاید ۱ از ۹ • زمان: ۰:۰۰ تا ۰:۴۵")
    add_notes(s1, "زمان هدف: ۰:۰۰ تا ۰:۴۵ (۴۵ ثانیه)\nسلام و درود. امروز پروژه اسلیپ‌لنز (SleepLens) را خدمت شما معرفی می‌کنیم؛ یک پلتفرم جامع هوش مصنوعی برای ارزیابی بالینی خواب. در پزشکی خواب، ثبت پلی‌سومنوگرافی استاندارد طلایی است، اما تحلیل آن ساعت‌ها زمان می‌برد و پیچیدگی مغز را به ۵ استیج گسسته تقلیل می‌دهد. ما با معماری دومرحله‌ای این چالش را حل کردیم: مرحله اول یک انسمبل دقیق برای استیجینگ ۳۰ ثانیه‌ای با تخمین اطمینان است، و مرحله دوم یک شبکه ترنسفورمر عمق پیوسته خواب همراه با استخراج ۹۶ پارامتر بالینی است که در قالب یک وب‌سایت در اختیار پزشک قرار می‌گیرد.")

    # =========================================================================
    # SLIDE 2: Dataset, Preprocessing & Standardization (FA)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "مجموعه داده، پیش‌پردازش و استانداردسازی", "اعتبار بنچ‌مارک و اصول تجربی",
               "پایگاه Sleep-EDF Expanded: ۱۹۷ شب کامل با اعتبارسنجی متقاطع بدون نشت اطلاعات (Subject-Atomic)", "۰:۴۵ – ۱:۳۰ | ۴۵s")

    # Right Card: Dataset
    add_card(s2, Inches(6.85), Inches(1.6), Inches(5.68), Inches(5.1))
    tb_ds = s2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.18), Inches(4.7))
    tf_ds = tb_ds.text_frame
    tf_ds.word_wrap = True
    tf_ds.margin_left = tf_ds.margin_top = tf_ds.margin_right = tf_ds.margin_bottom = 0

    p = tf_ds.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "📊 بنچ‌مارک رسمی Sleep-EDF Expanded (PhysioNet)"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = tf_ds.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "مقیاس و حجم داده‌ها:"
    p.font.name = FONT_FA
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(12)
    p = tf_ds.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "• ۱۹۷ ثبت شبانه کامل (معادل حدود ۲,۰۰۰ ساعت سیگنال چندکاناله)\n• ۲۳۷,۹۳۶ ایپاک استاندارد ۳۰ ثانیه‌ای برچسب‌گذاری‌شده توسط متخصص"
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(3)

    p = tf_ds.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "تفکیک دقیق دو کوهورت بدون تداخل:"
    p.font.name = FONT_FA
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(12)
    p = tf_ds.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "• کوهورت کاست (SC): ۱۵۳ ثبت افراد سالم در خانه (سیگنال هولتر روزمره)\n• کوهورت تله‌متری (ST): ۴۴ ثبت بالینی داخل بیمارستان و کلینیک خواب"
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(3)

    p = tf_ds.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "پروتکل اعتبارسنجی بدون نشت اطلاعات (Subject-Atomic 5-Fold CV):"
    p.font.name = FONT_FA
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(12)
    p = tf_ds.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "• تفکیک کاملا سوژه‌محور (folds5_v3)؛ هیچ بیماری به طور مشترک در بخش آموزش و ارزیابی قرار ندارد تا از حفظ‌کردن داده‌ها جلوگیری شود."
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(3)

    # Left Card: Preprocessing
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.65), Inches(5.1))
    tb_pr = s2.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.15), Inches(4.7))
    tf_pr = tb_pr.text_frame
    tf_pr.word_wrap = True
    tf_pr.margin_left = tf_pr.margin_top = tf_pr.margin_right = tf_pr.margin_bottom = 0

    p = tf_pr.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "⚙️ خط لوله پیش‌پردازش و یکپارچه‌سازی سیگنال‌ها"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    steps_fa = [
        ("۱. بازنمونه‌برداری استاندارد ۱۰۰ هرتز (Canonical 100 Hz):", "اعمال فیلتر ضددایره‌ای (Anti-aliasing) چندفازی روی کلیه کانال‌های خام به منظور رفع تفاوت نرخ نمونه‌برداری دستگاه‌های مختلف."),
        ("۲. استانداردسازی و نگاشت الکترودها:", "همگن‌سازی کانال‌های ناهمگون به الکترودهای استاندارد: مغز پیشانی (Fpz-Cz)، مغز پس‌سری (Pz-Oz)، حرکت چشم (EOG horiz) و عضله چانه (EMG)."),
        ("۳. پنجره‌بندی خاموشی چراغ‌ها (BENCHMARK_30):", "برش خودکار دوره‌های بیداری طولانی قبل و بعد از خواب ضمن حفظ کامل گذارهای فیزیولوژیک شروع و پایان خواب.")
    ]

    for title, desc in steps_fa:
        p = tf_pr.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = title
        p.font.name = FONT_FA
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)
        p = tf_pr.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = desc
        p.font.name = FONT_FA
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(3)

    add_footer(s2, "پروتکل اعتبارسنجی: ۵ فولد تفکیک‌شده بر اساس سوژه با پنجره استاندارد BENCHMARK_30", "تعداد کل ایپاک‌ها: ۲۳۷,۹۳۶ ایپاک معتبر")
    add_notes(s2, "زمان هدف: ۰:۴۵ تا ۱:۳۰ (۴۵ ثانیه)\nبرای اطمینان از اعتبار علمی، ما بنچ‌مارک کامل Sleep-EDF Expanded را با ۱۹۷ شب و ۲۳۸ هزار ایپاک به کار بردیم. برخلاف بسیاری از مقالات که دچار نشت سوژه بین فولدها هستند، ما فولدبندی کاملا سوژه‌محور انجام دادیم. هر دو کوهورت خانگی و کلینیکی به طور جداگانه تست شدند. تمامی سیگنال‌ها به ۱۰۰ هرتز استاندارد رسانده شدند و با پنجره‌بندی هوشمند، داده‌های غیرواقعی بیداری قبل از خواب حذف شدند.")

    # =========================================================================
    # SLIDE 3: Tier 1 Model — SSC Benchmarks (FA)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "مدل اول: دسته‌بندی ۳۰ ثانیه‌ای مراحل خواب (SSC)", "مدل مرحله اول هوش مصنوعی",
               "ارزیابی ساختاریافته بیش از ۲۵ پیکربندی روی تمامی ۱۹۷ شب داده", "۱:۳۰ – ۲:۲۵ | ۵۵s")

    # Table
    rows, cols = 7, 10
    tbl_shape = s3.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.733), Inches(3.1))
    table = tbl_shape.table

    col_widths = [Inches(1.1), Inches(3.633), Inches(0.95), Inches(0.85), Inches(0.85), Inches(0.9), Inches(0.85), Inches(0.85), Inches(0.85), Inches(0.9)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    headers_fa = ["کد اجرا", "معماری مدل و چیدمان سنسورها", "F1 کل", "SC F1", "ST F1", "N1 F1", "N2 F1", "N3 F1", "REM F1", "کاپا κ"]
    data_fa = [
        ["E19 (برتر)", "انسمبل ۴ مدله متنوع (ANY3ch + ANY2ch + USleep + LGBM)", "0.8425", "0.836", "0.865", "0.635", "0.892", "0.833", "0.915", "0.831"],
        ["E11", "مدل AnySleep ۳ کاناله (Fpz+Pz+EOG) — بهترین تک‌مدل", "0.8267", "0.809", "0.878", "0.622", "0.880", "0.788", "0.904", "0.812"],
        ["E11d", "مدل AnySleep ۲ سنسوره (Fpz+EOG) — بهترین ۲ سنسوره", "0.8227", "0.811", "0.865", "0.592", "0.880", "0.812", "0.905", "0.808"],
        ["E20", "مدل RobustSleepNet فرونتال (پاک — بدون پیش‌آموزش)", "0.7501", "0.736", "0.788", "0.442", "0.828", "0.717", "0.855", "0.724"],
        ["E00", "مدل LightGBM (۱۱ ویژگی طیفی و توان فرکانسی)", "0.6672", "—", "—", "0.339", "0.752", "0.687", "0.585", "0.612"],
        ["E01", "مدل سنتی بالینی YASA (جنگل تصادفی Fpz-Cz)", "0.5225", "0.486", "0.611", "0.100", "0.737", "0.538", "0.557", "0.473"],
    ]

    for c, h in enumerate(headers_fa):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_FA
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT

    for r, row in enumerate(data_fa):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = WIN_HIGHLIGHT_BG
            else:
                cell.fill.fore_color.rgb = INNER_CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_FA
            p.font.size = Pt(10.5)
            p.alignment = PP_ALIGN.RIGHT if c == 1 else PP_ALIGN.CENTER
            if r == 0:
                p.font.bold = True
                p.font.color.rgb = TEAL_PRIMARY if c in [0, 2, 5, 8, 9] else TEXT_WHITE
            else:
                p.font.color.rgb = TEXT_WHITE if c in [0, 1] else (ROSE_ACCENT if (r==5 and c in [2,5]) else TEXT_MUTED)

    # 3 Insights cards RTL order: Right, Middle, Left
    insights_fa = [
        ("۱. برتری تنوع ویژگی بر مقیاس", "افزودن مدل سبک LightGBM (با F1=0.667) به شبکه‌های عمیق باعث افزایش 0.016+ در دقت شد، چرا که ویژگی‌های طیفی مستقل از شکل موج خام را لحاظ می‌کند.", TEAL_PRIMARY, Inches(8.88)),
        ("۲. رکوردشکنی در تفکیک استیج‌ها", "سیستم برنده E19 بالاترین دقت تاریخ این بنچ‌مارک را در ۴ استیج اصلی ثبت کرد: N1 (0.635), N2 (0.892), N3 (0.833) و REM (0.915).", CYAN_ACCENT, Inches(4.84)),
        ("۳. رد الگوریتم ویتربی", "اعمال دیکودینگ ویتربی بر احتمالات گذار باعث افت دقت شد (0.8267 به 0.8259)، چرا که معماری AnySleep ذاتاً میدان دید زمانی ۱۴ دقیقه‌ای دارد.", AMBER_WARN, Inches(0.8)),
    ]

    for title, text, col, left_x in insights_fa:
        add_card(s3, left_x, Inches(4.88), Inches(3.64), Inches(1.85))
        tb = s3.shapes.add_textbox(left_x + Inches(0.15), Inches(4.98), Inches(3.34), Inches(1.65))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = title
        p.font.name = FONT_FA
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.RIGHT
        p2.text = text
        p2.font.name = FONT_FA
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

    add_footer(s3, "انسمبل برنده ۴ مدله: توافق مستقل در تمام ۵ فولد ارزیابی OOF", "ضریب کاپای کوهن: 0.831 (معادل توافق پزشکان فوق تخصص)")
    add_notes(s3, "زمان هدف: ۱:۳۰ تا ۲:۲۵ (۵۵ ثانیه)\nبرای دسته‌بندی استیج‌های خواب، ما ۲۵ پیکربندی مختلف را بنچ‌مارک کردیم. مدل‌های سنتی مثل YASA در استیج دشوار N1 تنها به دقت ۱۰ درصد می‌رسند. مدل‌های یادگیری عمیق مانند AnySleep عملکرد فوق‌العاده 0.8267 را ثبت کردند. اما سیستم نهایی ما، انسمبل E19 با ترکیب AnySleep و U-Sleep و مدل آماری LightGBM به رکورد F1 ماکرو 0.8425 و ضریب کاپای 0.831 رسید. جالب اینجاست که مدل آماری LightGBM به دلیل تفاوت در فضای ویژگی توانست خطاهای مدل‌های عمیق را به خوبی جبران کند.")

    # =========================================================================
    # SLIDE 4: Sensor Ablation & Explainability (FA)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "کاهش سنسورها و هوش مصنوعی تفسیرپذیر (XAI)", "تحلیل‌های فنی استیجینگ",
               "کشف حداقل چیدمان سخت‌افزاری پوشیدنی و بازرسی توجه فیزیولوژیک شبکه عمیق", "۲:۲۵ – ۳:۱۵ | ۵۰s")

    # Right: Channel Ablation
    add_card(s4, Inches(6.85), Inches(1.6), Inches(5.68), Inches(5.1))
    tb_ab = s4.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.18), Inches(4.7))
    tf_ab = tb_ab.text_frame
    tf_ab.word_wrap = True
    tf_ab.margin_left = tf_ab.margin_top = tf_ab.margin_right = tf_ab.margin_bottom = 0

    p = tf_ab.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "🔌 ارزیابی کاهش کانال‌ها برای گجت‌های پوشیدنی"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    ablation_fa = [
        ("سیستم ۳ کاناله کامل (Fpz-Cz + Pz-Oz + EOG):", "0.8267 F1 (مرجع)", False),
        ("★ چیدمان ۲ سنسوره (Fpz-Cz + EOG):", "0.8227 F1 (تنها 0.004- افت!)", True),
        ("دو کانال مغزی بدون چشم (Fpz-Cz + Pz-Oz):", "0.8134 F1 (افت 0.013)", False),
        ("تک الکترود پیشانی (Fpz-Cz):", "0.8015 F1 (افت 0.025)", False),
        ("تک الکترود پس‌سری (Pz-Oz):", "0.7736 F1 (افت 0.053)", False),
    ]

    for label, score, highlight in ablation_fa:
        p = tf_ab.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"{label}  ◀──  {score}"
        p.font.name = FONT_FA
        p.font.size = Pt(12.5)
        p.space_before = Pt(8)
        if highlight:
            p.font.bold = True
            p.font.color.rgb = CYAN_ACCENT
        else:
            p.font.color.rgb = TEXT_WHITE if "کامل" in label else TEXT_MUTED

    p = tf_ab.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "نتیجه بالینی: کانال پس‌سری Pz-Oz کاملا قابل حذف است. افزودن EOG به الکترود پیشانی به طور مستقیم سیگنال دو‌قطبی چشم را ثبت کرده و ریکال خواب REM را به 0.913 می‌رساند (بالاتر از 0.897 در سیستم ۳ کاناله!)."
    p.font.name = FONT_FA
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(14)

    # Left: Attention & Confidence
    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.65), Inches(5.1))
    tb_att = s4.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.15), Inches(4.7))
    tf_att = tb_att.text_frame
    tf_att.word_wrap = True
    tf_att.margin_left = tf_att.margin_top = tf_att.margin_right = tf_att.margin_bottom = 0

    p = tf_att.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "🔍 فیزیولوژی توجه لایه‌ای و کالیبراسیون اطمینان"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_att.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "توزیع توجه لایه ۱۲ شبکه (Skip Connections) کاملا منطبق بر قواعد پزشکی:"
    p.font.name = FONT_FA
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(10)

    attn_fa = [
        ("• بیداری (Wake):", "۶۷.۲٪ توجه روی Pz", "(ردیابی محو شدن ریتم آلفا)"),
        ("• خواب عمیق (N3):", "۹۳.۰٪ توجه روی امواج مغزی", "(امواج دلتای کورتکس، چشم آرام)"),
        ("• خواب رویابینی (REM):", "۵۲.۰٪ توجه روی EOG", "(ردیابی پرش‌های سریع چشم در آتونی)"),
    ]
    for stage, stat, note in attn_fa:
        p = tf_att.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"{stage} {stat} {note}"
        p.font.name = FONT_FA
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(5)

    p = tf_att.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "کالیبراسیون یکنوای اطمینان مدل (ECE = 0.029):"
    p.font.name = FONT_FA
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(12)

    p = tf_att.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "• اطمینان بالای ۹۰٪ (۳۰.۳٪ کل شب): دارای دقت تجربی ۹۹.۵٪!\n• اطمینان زیر ۶۰٪ (~۱۵٪ کل شب): برچسب‌گذاری با needs_review = true\n• پزشک در چرخه (Doctor-in-the-Loop): متخصص فقط ایپاک‌های مرزی را چک می‌کند."
    p.font.name = FONT_FA
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(5)

    add_footer(s4, "هدف سخت‌افزاری: هدبند خانگی با ۲ سنسور پیشانی (Fpz-Cz + EOG)", "خطای کالیبراسیون مورد انتظار: ECE = 0.029")
    add_notes(s4, "زمان هدف: ۲:۲۵ تا ۳:۱۵ (۵۰ ثانیه)\nدو سوال عملی مهم: آیا می‌توان این سیستم را با یک گجت پوشیدنی ساده اجرا کرد و آیا هوش مصنوعی ما قابل تفسیر است؟ نتایج تحلیل حذف سنسورها نشان داد با فقط ۲ سنسور—یک الکترود پیشانی و یک سنسور چشم—می‌توان به دقت 0.8227 رسید که فقط 0.004 با سیستم کامل فاصله دارد. همچنین لایه‌های توجه شبکه دقیقا رفتار پزشک را تقلید می‌کنند: در بیداری به ریتم آلفا در پشت سر نگاه می‌کنند و در خواب REM پنجاه و دو درصد توجه به سنسور چشم می‌رود. خروجی اطمینان مدل نیز به قدری دقیق کالیبره شده که ایپاک‌های با اطمینان بالای ۹۰ درصد، دقت ۹۹.۵ درصدی دارند.")

    # =========================================================================
    # SLIDE 5: Tier 2 Model — Continuous Depth (SDI) & 96 Features (FA)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "مدل دوم: عمق پیوسته خواب (SDI) و موتور ۹۶ ویژگی", "مدل مرحله دوم هوش مصنوعی",
               "گذار از ۵ پله گسسته به عمق پیوسته همگام‌سازی بیولوژیکی مغز", "۳:۱۵ – ۴:۰۵ | ۵۰s")

    # Top Card: SDI Transformer
    add_card(s5, Inches(0.8), Inches(1.6), Inches(11.733), Inches(2.35), border=CYAN_ACCENT, border_width=2)
    tb_sdi = s5.shapes.add_textbox(Inches(1.05), Inches(1.75), Inches(11.233), Inches(2.05))
    tf_sdi = tb_sdi.text_frame
    tf_sdi.word_wrap = True
    tf_sdi.margin_left = tf_sdi.margin_top = tf_sdi.margin_right = tf_sdi.margin_bottom = 0

    p = tf_sdi.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "شبکه ترنسفورمر SDI (منتشر شده در مجله معتبر npj Digital Medicine 2025)"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_sdi.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "تخمین شاخص پیوسته عمق خواب (۰.۰ = بیداری/خواب سبک تا ۱.۰ = اوج امواج آهسته عمیق) برای هر ایپاک ۳۰ ثانیه‌ای. برخلاف استیجینگ گسسته سنتی، مدل SDI می‌تواند میکرو-بیداری‌ها، تخلیه بار خستگی روزانه و الگوهای متناوب حلقوی (CAP) را به دقت ردیابی کند."
    p.font.name = FONT_FA
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(5)

    p = tf_sdi.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "بردار پنج‌گانه استاندارد تحقیقاتی ژو و همکاران:  [RB: بار خواب سطحی SDI < 0.20]  •  [AP: میانگین عمق خواب]  •  [CV: ناپایداری و نوسان عمق]  •  [MDR: میانگین عمق در REM]  •  [PR: سهم فاز REM]"
    p.font.name = FONT_FA
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_before = Pt(8)

    # Bottom Card: 96-Feature Engine
    add_card(s5, Inches(0.8), Inches(4.1), Inches(11.733), Inches(2.65))
    tb_feat = s5.shapes.add_textbox(Inches(1.05), Inches(4.25), Inches(11.233), Inches(2.35))
    tf_feat = tb_feat.text_frame
    tf_feat.word_wrap = True
    tf_feat.margin_left = tf_feat.margin_top = tf_feat.margin_right = tf_feat.margin_bottom = 0

    p = tf_feat.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "📈 موتور استخراج پارامترهای جامع فیزیولوژیک (~۹۶ شاخص بالینی)"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    # 4 columns RTL
    f_cols_fa = [
        ("پیوستگی و پراکندگی خواب", "• زمان کل خواب، TIB، بهره‌وری SE%\n• تاخیر شروع خواب، WASO، تاخیر REM\n• شاخص بیداری‌های مکرر\n• شاخص قطعه‌قطعه شدن خواب (SFI)"),
        ("ساختار ماکرو و دینامیک", "• سهم درصدی استیج‌های N1, N2, N3, REM\n• جابجایی استیج‌ها در نیمه اول و دوم شب\n• طول دوره‌های خواب مداوم\n• ماتریس گذار مارکوف ۵×۵"),
        ("دینامیک طیفی مغز (EEG)", "• توان طیفی ولچ (دلتا، تتا، آلفا، سیگما، بتا)\n• فعالیت امواج آهسته مغزی (SWA)\n• آنتروپی طیفی سیگنال\n• آنتروپی جایگشتی در فازهای مختلف"),
        ("ریزساختار و سیگنال عضله", "• چگالی و دامنه دوک‌های خواب YASA (N2)\n• دامنه و مدت امواج آهسته مغز\n• توان RMS عضله زیر چانه (EMG)\n• نسبت آتونی عضلانی در فاز REM"),
    ]

    for i, (title, desc) in enumerate(f_cols_fa):
        # RTL index: 0 is on the far right (x = 1.05 + 3*2.85)
        c_left = Inches(1.05 + (3 - i) * 2.85)
        box = s5.shapes.add_textbox(c_left, Inches(4.65), Inches(2.7), Inches(1.95))
        btf = box.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0
        p = btf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = title
        p.font.name = FONT_FA
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p2 = btf.add_paragraph()
        p2.alignment = PP_ALIGN.RIGHT
        p2.text = desc
        p2.font.name = FONT_FA
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(4)

    add_footer(s5, "دیدگاه دوگانه: هیپنوگرام گسسته کلان + عمق پیوسته و ریزساختارهای خواب", "خروجی: ۹۶ ویژگی بالینی برای هر شب آزمایش")
    add_notes(s5, "زمان هدف: ۳:۱۵ تا ۴:۰۵ (۵۰ ثانیه)\nدر مدل مرحله دوم، ما از محدودیت هیپنوگرام ۵ مرحله‌ای عبور کردیم. ما شبکه ترنسفورمر SDI منتشر شده در ژورنال معتبر npj Digital Medicine 2025 را پیاده‌سازی کردیم. این مدل یک عدد پیوسته بین صفر و یک برای عمق خواب در هر ایپاک تولید می‌کند. همگام با این مدل، موتور استخراج فیزیولوژیک ما ۹۶ پارامتر بالینی را اندازه می‌گیرد: پیوستگی خواب، ماتریس گذار مارکوف بین مراحل، توان باندهای فرکانسی مغز، و نشانگرهای ریزساختاری مثل چگالی دوک‌های خواب و آتونی عضلانی در REM.")

    # =========================================================================
    # SLIDE 6: Sleep Quality Index & Multi-Modal Scoring (FA)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "شاخص کیفیت خواب و امتیازدهی چندوجهی", "هوش تشخیصی بالینی",
               "سنتز ۹۶ پارامتر فیزیولوژیک عینی، عمق پیوسته و ادراک ذهنی بیمار", "۴:۰۵ – ۴:۵۵ | ۵۰s")

    # Right Card: SleepLens Score
    add_card(s6, Inches(6.85), Inches(1.6), Inches(5.68), Inches(5.1))
    tb_sc = s6.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.18), Inches(4.7))
    tf_sc = tb_sc.text_frame
    tf_sc.word_wrap = True
    tf_sc.margin_left = tf_sc.margin_top = tf_sc.margin_right = tf_sc.margin_bottom = 0

    p = tf_sc.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "🎯 امتیاز یکپارچه و شفاف کیفیت خواب (۰ تا ۱۰۰)"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = tf_sc.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "فرمول شفاف و غیرجعبه‌سیاه مبتنی بر ۷ مولفه کلیدی:"
    p.font.name = FONT_FA
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(4)

    score_parts_fa = [
        ("• بهره‌وری خواب (SE%):", "حداکثر ۲۵ امتیاز", "(نسبت خواب مفید به کل زمان در بستر)"),
        ("• عمق پیوسته خواب (Mean SDI):", "حداکثر ۲۵ امتیاز", "(میزان اشباع امواج آهسته کورتکس)"),
        ("• پیوستگی و عدم بیداری:", "حداکثر ۲۰ امتیاز", "(جریمه تغییرات مکرر استیج و بیداری)"),
        ("• سهم خواب عمیق (N3%):", "حداکثر ۱۰ امتیاز", "(ترمیم جسمانی و بازسازی سلولی)"),
        ("• سهم خواب رویابینی (REM%):", "حداکثر ۱۰ امتیاز", "(تثبیت حافظه و تعادل شناختی)"),
        ("• تاخیر به خواب رفتن (SOL):", "حداکثر ۵ امتیاز", "(معیار ورود روان به خواب زیر ۳۰ دقیقه)"),
        ("• مقاومت در برابر پراکندگی:", "حداکثر ۵ امتیاز", "(کنترل ایپاک‌های خواب سطحی SDI < 0.20)"),
    ]

    for label, pts, note in score_parts_fa:
        p = tf_sc.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"{label} {pts}  {note}"
        p.font.name = FONT_FA
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(5)

    p = tf_sc.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "رتبه‌بندی بالینی:  ۸۵ تا ۱۰۰: عالی  |  ۷۰ تا ۸۴: خوب  |  ۵۵ تا ۶۹: متوسط  |  زیر ۵۵: ضعیف"
    p.font.name = FONT_FA
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_before = Pt(8)

    # Left Card: PSQI & Synthesis
    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.65), Inches(5.1))
    tb_ps = s6.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.15), Inches(4.7))
    tf_ps = tb_ps.text_frame
    tf_ps.word_wrap = True
    tf_ps.margin_left = tf_ps.margin_top = tf_ps.margin_right = tf_ps.margin_bottom = 0

    p = tf_ps.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "🩺 تلفیق داده‌های عینی و ادراک ذهنی بیمار (PSQI)"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_ps.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "یکپارچگی استاندارد با پرسشنامه پیتزبورگ (PSQI):"
    p.font.name = FONT_FA
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(10)
    p = tf_ps.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "شامل هر ۷ مولفه بالینی (کیفیت ذهنی، تاخیر خواب، مدت خواب، بازدهی خواب، اختلالات، مصرف دارو، و اختلال عملکرد روزانه) با امتیاز کل ۰ تا ۲۱."
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(3)

    p = tf_ps.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "تشخیص بی‌خوابی متناقض (Sleep State Misperception):"
    p.font.name = FONT_FA
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = AMBER_WARN
    p.space_before = Pt(12)
    p = tf_ps.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "سیستم با مقایسه PSQI و PSG، بیمارانی که حس بیداری و نارضایتی شدید دارند (PSQI > 12) ولی در آزمایش مغزی ساختار خواب کاملا سالم دارند را فورا شناسایی می‌کند."
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(3)

    p = tf_ps.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "سنتز خودکار گزارش پزشکی با هوش مصنوعی:"
    p.font.name = FONT_FA
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    p.space_before = Pt(12)
    p = tf_ps.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "مدل زبانی هوشمند (gpt-6-luna) بر اساس ۹۶ پارامتر و عمق پیوسته گزارش متنی ساختاریافته پزشکی تولید می‌کند که آماده تایید و امضای پزشک است."
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(3)

    add_footer(s6, "کاربرد بالینی: تبدیل ۹۶ پارامتر پیچیده مهندسی به یک شاخص قابل‌فهم برای بیمار و پزشک", "فرمولاسیون شفاف و کاملا قابل تفسیر")
    add_notes(s6, "زمان هدف: ۴:۰۵ تا ۴:۵۵ (۵۰ ثانیه)\nبرای تبدیل این حجم از متغیرها به یک گزارش قابل فهم، ما امتیاز جامع اسلیپ‌لنز را طراحی کردیم. این امتیاز یک جعبه سیاه مبهم نیست؛ بلکه یک شاخص ۱۰۰ امتیازی کاملا شفاف است: ۲۵٪ بهره‌وری خواب، ۲۵٪ عمق پیوسته، ۲۰٪ پیوستگی و پایداری، و مابقی مربوط به خواب عمیق N3، خواب REM، تاخیر خواب و عدم پراکندگی است. علاوه بر این، ما پرسشنامه استاندارد پیتزبورگ (PSQI) را در سامانه تعبیه کردیم تا با تطبیق حس ذهنی بیمار و سیگنال‌های عینی، اختلالاتی مثل بی‌خوابی متناقض یا حس بیداری کاذب به سرعت کشف شوند.")

    # =========================================================================
    # SLIDE 7: End-to-End System Architecture (FA)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "معماری جامع سیستم و جریان پردازش داده‌ها", "طراحی نرم‌افزار و معماری سامانه",
               "خط لوله پردازش ناهمگام و مقیاس‌پذیر از آپلود فایل EDF تا گزارش تشخیصی نهایی", "۴:۵۵ – ۵:۴۵ | ۵۰s")

    # 4 Steps RTL: 01 on the right (x=9.8), 04 on the left (x=0.8)
    step_boxes_fa = [
        ("۰۱ • آپلود و امنیت", "• درگ‌اند‌دراپ EDF (تا ۵۰۰ مگابایت)\n• فرم اختیاری ۷‌بخشی PSQI\n• اتصال به پرونده بالینی بیمار\n• احراز هویت امن JWT با کوکی", TEAL_PRIMARY, Inches(9.8)),
        ("۰۲ • آماده‌سازی سیگنال", "• بازنمونه‌برداری ۱۰۰ هرتز\n• پنجره‌بندی خاموشی چراغ‌ها\n• کنترل کیفیت و نگاشت کانال‌ها\n• صف کارگران ناهمگام Celery", CYAN_ACCENT, Inches(6.8)),
        ("۰۳ • موتورهای دوگانه AI", "• مدل اول: انسمبل ۴ مدله\n  (استیجینگ ۳۰ثانیه + اطمینان)\n• مدل دوم: ترنسفورمر SDI\n  (منحنی عمق پیوسته ۰ تا ۱)", RGBColor(56, 189, 248), Inches(3.8)),
        ("۰۴ • پلتفرم و هوش نهایی", "• امتیاز ۰ تا ۱۰۰ اسلیپ‌لنز\n• موتور ۹۶ ویژگی بالینی\n• گزارش خودکار متنی با LLM\n• چت‌بات تشخیصی OpenCode", EMERALD_GREEN, Inches(0.8)),
    ]

    for title, desc, color, left_pos in step_boxes_fa:
        add_card(s7, left_pos, Inches(1.6), Inches(2.733), Inches(3.1), border=color, border_width=2)
        tb = s7.shapes.add_textbox(left_pos + Inches(0.12), Inches(1.8), Inches(2.493), Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = title
        p.font.name = FONT_FA
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.RIGHT
        p2.text = desc
        p2.font.name = FONT_FA
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(8)

    # Technology Stack Box
    add_card(s7, Inches(0.8), Inches(4.9), Inches(11.733), Inches(1.85))
    tb_st = s7.shapes.add_textbox(Inches(1.05), Inches(5.05), Inches(11.233), Inches(1.55))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True
    tf_st.margin_left = tf_st.margin_top = tf_st.margin_right = tf_st.margin_bottom = 0
    p = tf_st.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "استک فناوری و ارتباطات سیستم در بستر نرم‌افزاری کاملاً کانتینریزه (Docker & Compose):"
    p.font.name = FONT_FA
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p = tf_st.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "• فرانت‌اند: Next.js 14 App Router, TypeScript, Tailwind CSS, Zustand, React Query.\n• هسته بک‌اند: Django REST Framework, PostgreSQL, کارگران پردازش ناهمگام Celery, پیام‌رسان Redis.\n• هسته یادگیری ماشین: PyTorch, MNE-Python, LightGBM, YASA, پردازش سیگنال SciPy.\n• لایه استدلال و هوش مصنوعی: OpenAI API (gpt-6-luna) برای تولید گزارش تشخیصی؛ ایجنت محلی OpenCode برای مشاوره پزشکی."
    p.font.name = FONT_FA
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(6)

    add_footer(s7, "معماری مستقل و نامتقارن: پردازش سنگین مدل‌ها هرگز درخواست‌های وب را مسدود نمی‌کند", "میانگین زمان پردازش کامل هر شب: کمتر از ۳۵ ثانیه روی CPU معمولی")
    add_notes(s7, "زمان هدف: ۴:۵۵ تا ۵:۴۵ (۵۰ ثانیه)\nدر این اسلاید معماری مهندسی سامانه اسلیپ‌لنز را مشاهده می‌کنید. پزشک یا اپراتور فایل خام EDF را از طریق پنل مدرن Next.js آپلود می‌کند. بک‌اند جانگو درخواست را اعتبارسنجی کرده و پردازش سنگین را به کارگران Celery با صف ردیس می‌سپارد. سیگنال‌ها همگن‌سازی شده و همزمان به دو مدل هدایت می‌شوند: انسمبل استیجینگ و ترنسفورمر عمق خواب. داده‌ها همراه با ۹۶ ویژگی در پایگاه داده ذخیره شده و مدل زبانی هوشمند گزارش تشخیصی را تدوین می‌کند، در حالی که ایجنت محلی OpenCode به سوالات پزشک پاسخ می‌دهد.")

    # =========================================================================
    # SLIDE 8: Clinical Web Application Walkthrough (FA)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "نمای وب‌اپلیکیشن بالینی (SleepLens UI/UX)", "محصول نهایی و پنل پزشک",
               "کابین خلبان تشخیصی جامع طراحی‌شده اختصاصی برای پزشکان و متخصصان تکنولوژی خواب", "۵:۴۵ – ۶:۳۰ | ۴۵s")

    ui_features_fa = [
        ("🎯 کارت امتیاز اسلیپ‌لنز و شاخص‌های کلیدی (KPIs)",
         "• گیج دایره‌ای امتیاز ۰ تا ۱۰۰ با تفکیک نموداری ۷ مولفه فیزیولوژیک.\n• نمایش شاخص‌های اصلی: زمان خواب (۴۱۲.۵ دقیقه)، بهره‌وری خواب (۸۸.۴٪)، تاخیر خواب (۱۴.۰ دقیقه)، بیداری پس از خواب (۳۸.۵ دقیقه)، تاخیر REM (۷۸.۰ دقیقه)، و میانگین اطمینان مدل (۹۱.۲٪).",
         TEAL_PRIMARY, Inches(6.85), Inches(1.6)),
        ("📊 هیپنوگرام تعاملی با لایه اطمینان مدل",
         "• تایم‌لاین همگام‌شده مراحل خواب (Wake, REM, N1, N2, N3).\n• لایه بصری اطمینان مدل: بالا (≥۸۰٪)، متوسط (≥۶۰٪)، پایین (<۶۰٪).\n• نشانگرهای پرچم بازبینی با یک کلیک برای ایپاک‌های مرزی با اطمینان پایین.",
         CYAN_ACCENT, Inches(0.8), Inches(1.6)),
        ("📈 منحنی عمق پیوسته SDI و کاوشگر سیگنال خام",
         "• شکل‌موج پیوسته ۰ تا ۱ عمق خواب: نمایش چرخه‌های عمیق امواج آهسته مغز.\n• کاوشگر بدون تاخیر سیگنال: بررسی آنی سیگنال‌های میکروولتی خام برای هر ایپاک ۳۰ ثانیه‌ای دلخواه (Fpz-Cz, Pz-Oz, EOG, EMG).",
         RGBColor(56, 189, 248), Inches(6.85), Inches(4.25)),
        ("💬 گزارش متنی خودکار و چت‌بات مشاور بالینی",
         "• گزارش تشخیصی با LLM: گزارش ساختاریافته مارک‌داون با امکان ویرایش و تایید پزشک.\n• دستیار گفتگوی OpenCode: چت‌بات درون‌جلسه‌ای برای پاسخ به سوالات بالینی پزشک با اتکا به دیتای دقیق بیمار.",
         EMERALD_GREEN, Inches(0.8), Inches(4.25)),
    ]

    for title, desc, col, left_x, top_y in ui_features_fa:
        add_card(s8, left_x, top_y, Inches(5.68), Inches(2.45), border=col, border_width=2)
        tb = s8.shapes.add_textbox(left_x + Inches(0.18), top_y + Inches(0.18), Inches(5.32), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = title
        p.font.name = FONT_FA
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.RIGHT
        p2.text = desc
        p2.font.name = FONT_FA
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(8)

    add_footer(s8, "فرانت‌اند واکنش‌گرای Next.js • رابط کاربری پیشرفته بالینی با تم مدرن Slate & Teal", "کاهش بیش از ۸۵ درصدی زمان بررسی پرونده توسط پزشک")
    add_notes(s8, "زمان هدف: ۵:۴۵ تا ۶:۳۰ (۴۵ ثانیه)\nدر این اسلاید رابط کاربری اسلیپ‌لنز را می‌بینید. در داشبورد، پزشک امتیاز کیفی خواب و شاخص‌های کلیدی را در یک نگاه بررسی می‌کند. هیپنوگرام تعاملی به همراه میزان اطمینان مدل نمایش داده می‌شود و ایپاک‌های مبهم با پرچم زرد مشخص شده‌اند تا پزشک بتواند با یک کلیک آن‌ها را بازبینی کند. همچنین منحنی پیوسته عمق خواب، کاوشگر سیگنال‌های خام میکروولتی مغز برای هر ۳۰ ثانیه، پیش‌نویس گزارش متنی، و چت‌بات دستیار تشخیصی، زمان اسکورینگ را از ۹۰ دقیقه به چند دقیقه کاهش می‌دهند.")

    # =========================================================================
    # SLIDE 9: Deployment Performance & Roadmap (FA)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "عملکرد اجرایی، تاثیر بالینی و نقشه راه آینده", "جمع‌بندی و افق توسعه",
               "بنچ‌مارک‌های سرعت در دنیای واقعی، ارتقای بهره‌وری درمان و ترجمه به گجت‌های پوشیدنی", "۶:۳۰ – ۷:۰۰ | ۳۰s")

    # Right: Speed Benchmarks
    add_card(s9, Inches(6.85), Inches(1.6), Inches(5.68), Inches(5.1))
    tb_sp = s9.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.18), Inches(4.7))
    tf_sp = tb_sp.text_frame
    tf_sp.word_wrap = True
    tf_sp.margin_left = tf_sp.margin_top = tf_sp.margin_right = tf_sp.margin_bottom = 0

    p = tf_sp.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "⚡ بنچ‌مارک سرعت روی پردازنده معمولی (CPU بدون نیاز به GPU)"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    speed_rows_fa = [
        ("• استنتاج هر ایپاک ۳۰ ثانیه‌ای:", "۳.۷ تا ۴.۷ میلی‌ثانیه", "(۶,۴۰۰ برابر سریع‌تر از زمان واقعی)"),
        ("• استنتاج یک شب کامل ۲۲ ساعته (۲,۶۵۰ ایپاک):", "۱۲.۵ ثانیه", "(انسمبل استیجینگ ۴ مدله)"),
        ("• استنتاج عمق پیوسته ترنسفورمر SDI:", "۸.۲ ثانیه", "(شبکه توجه عمیق)"),
        ("• استخراج ۹۶ پارامتر فیزیولوژیک:", "۹.۸ ثانیه", "(تحلیل‌های ولچ و دوک‌های خواب YASA)"),
    ]

    for label, val, note in speed_rows_fa:
        p = tf_sp.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"{label} {val} {note}"
        p.font.name = FONT_FA
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    p = tf_sp.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "★ زمان کل پردازش شبانه: کمتر از ۳۵ ثانیه برای کل شب!"
    p.font.name = FONT_FA
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    p.space_before = Pt(16)

    p = tf_sp.add_paragraph()
    p.alignment = PP_ALIGN.RIGHT
    p.text = "تاثیر بر بهره‌وری بالینی:\nکاهش بیش از ۸۵ درصدی زمان اسکورینگ چشمی توسط پزشک همراه با کشف ناهنجاری‌های پنهان مغزی و ردپای دیجیتال تصمیم‌گیری."
    p.font.name = FONT_FA
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(10)

    # Left: Roadmap
    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.65), Inches(5.1))
    tb_rd = s9.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.15), Inches(4.7))
    tf_rd = tb_rd.text_frame
    tf_rd.word_wrap = True
    tf_rd.margin_left = tf_rd.margin_top = tf_rd.margin_right = tf_rd.margin_bottom = 0

    p = tf_rd.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "🚀 نقشه راه ترجمه بالینی و توسعه آینده"
    p.font.name = FONT_FA
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    roadmap_fa = [
        ("۱. به‌کارگیری گجت‌های پوشیدنی روزمره:", "پیاده‌سازی مدل‌های جانشین آموزش‌دیده (ResUNet و WatchSleepNet) برای ساعت‌های هوشمند مبتنی بر اکتی‌گرافی و سنسور PPG بدون نیاز به اتصالات مغزی."),
        ("۲. اسکورینگ تخصصی آپنه و ناهنجاری حرکتی:", "ارتقای آشکارسازها به استانداردهای رسمی پزشکی برای سنجش شاخص آپنه-هایپوپنه (AHI) و سندروم پای بی‌قرار (PLMD)."),
        ("۳. کارآزمایی بالینی در مراکز متعدد:", "اعتبارسنجی روی دیتابیس‌های چندمرکزی بین‌المللی (SHHS) و گروه‌های مبتلا به بیماری‌های نورولوژیک (نظیر اختلال رفتاری خواب REM در پارکینسون).")
    ]

    for title, desc in roadmap_fa:
        p = tf_rd.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = title
        p.font.name = FONT_FA
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(14)
        p = tf_rd.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.text = desc
        p.font.name = FONT_FA
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(3)

    add_footer(s9, "اسلیپ‌لنز: ارائه دقت در سطح استاندارد طلایی با سرعت پردازش بلادرنگ", "از توجه و همراهی شما سپاسگزاریم • پرسش و پاسخ")
    add_notes(s9, "زمان هدف: ۶:۳۰ تا ۷:۰۰ (۳۰ ثانیه)\nدر جمع‌بندی، اسلیپ‌لنز با کسب نمره F1 معادل 0.8425 و کاپای 0.831 در استیجینگ، و پردازش کامل یک شب ظرف کمتر از ۳۵ ثانیه روی پردازنده‌های معمولی، پلی میان هوش مصنوعی پیشرفته و نیازهای روزمره کلینیک‌های خواب ایجاد کرده است. همچنین اثبات کارایی چیدمان ۲ سنسوره راه را برای هدبندهای خانگی هموار می‌کند. اسلیپ‌لنز زمان پزشک را ۸۵ درصد کاهش داده و کیفیت مراقبت از بیمار را ارتقا می‌بخشد. از توجه شما بسیار متشکرم.")

    output_path = "SleepLens_Presentation_FA.pptx"
    prs.save(output_path)
    print(f"Persian presentation saved successfully to {os.path.abspath(output_path)}")

if __name__ == "__main__":
    build_persian_deck()
