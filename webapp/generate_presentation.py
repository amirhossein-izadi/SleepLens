import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette - Professional Medical AI Theme
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    CARD_BG = RGBColor(30, 41, 59)       # Slate 800
    BORDER_COLOR = RGBColor(51, 65, 85)  # Slate 700
    TEXT_WHITE = RGBColor(248, 250, 252) # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)# Slate 400
    TEAL_PRIMARY = RGBColor(20, 184, 166)# Teal 500
    CYAN_ACCENT = RGBColor(6, 182, 212)  # Cyan 500
    EMERALD_GREEN = RGBColor(16, 185, 129)# Green
    AMBER_WARN = RGBColor(245, 158, 11)  # Amber
    BLUE_LIGHT = RGBColor(56, 189, 248)  # Sky 400

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="SLEEPLENS • CLINICAL AI PLATFORM", timing_text=""):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = TEAL_PRIMARY
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_before = Pt(4)

        if timing_text:
            time_box = slide.shapes.add_textbox(Inches(9.5), Inches(0.4), Inches(3.0), Inches(0.5))
            ttf = time_box.text_frame
            ttf.word_wrap = True
            tp = ttf.paragraphs[0]
            tp.alignment = PP_ALIGN.RIGHT
            tp.text = timing_text
            tp.font.size = Pt(11)
            tp.font.bold = True
            tp.font.color.rgb = CYAN_ACCENT

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    def add_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # =========================================================================
    # SLIDE 1: Title & Executive Summary
    # =========================================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)
    
    # Hero Title Box
    h_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(2.2))
    htf = h_box.text_frame
    htf.word_wrap = True
    
    cat_p = htf.paragraphs[0]
    cat_p.text = "CLINICAL AI & NEUROPHYSIOLOGY"
    cat_p.font.size = Pt(12)
    cat_p.font.bold = True
    cat_p.font.color.rgb = TEAL_PRIMARY
    
    main_p = htf.add_paragraph()
    main_p.text = "SleepLens: Dual-Tier AI for Sleep Stage Classification\n& Continuous Sleep Quality Assessment"
    main_p.font.size = Pt(30)
    main_p.font.bold = True
    main_p.font.color.rgb = TEXT_WHITE
    main_p.space_before = Pt(8)

    sub_p = htf.add_paragraph()
    sub_p.text = "Bridging 30-Second Discrete Staging with Continuous Depth Dynamics & Microstructural Phenotyping"
    sub_p.font.size = Pt(14)
    sub_p.font.color.rgb = CYAN_ACCENT
    sub_p.space_before = Pt(6)

    # 3 Summary Cards
    c1 = add_card(s1, Inches(0.8), Inches(3.6), Inches(3.64), Inches(3.2))
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(3.24), Inches(2.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "01 • The Problem"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AMBER_WARN
    p2 = tf1.add_paragraph()
    p2.text = "• Manual PSG takes 1.5–2 hours/night.\n• Coarse 5-stage AASM ignores micro-arousals and continuous depth.\n• High inter-scorer variance on N1 (<60%) and boundary transitions."
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(8)

    c2 = add_card(s1, Inches(4.84), Inches(3.6), Inches(3.64), Inches(3.2))
    tb2 = s1.shapes.add_textbox(Inches(5.04), Inches(3.8), Inches(3.24), Inches(2.8))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "02 • Dual-Model AI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p2 = tf2.add_paragraph()
    p2.text = "• Tier 1: 4-Model Ensemble Macro-F1 0.8425 (Cohen's κ 0.831) on 237k epochs.\n• Tier 2: SDI Transformer (npj Digital Med) continuous depth (0-1) + 96 clinical parameters."
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(8)

    c3 = add_card(s1, Inches(8.88), Inches(3.6), Inches(3.64), Inches(3.2))
    tb3 = s1.shapes.add_textbox(Inches(9.08), Inches(3.8), Inches(3.24), Inches(2.8))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "03 • The Product"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    p2 = tf3.add_paragraph()
    p2.text = "• Next.js + Django REST clinical webapp.\n• Interactive hypnograms, SDI depth curves, and raw signal inspection.\n• Automated LLM reports + OpenCode clinical consultation chat."
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(8)

    add_notes(s1, "Target: 0:00 - 0:45 (45s)\nGood morning. Today we present SleepLens, an end-to-end clinical AI platform designed to transform sleep diagnostics. Polysomnography is gold standard, yet analyzing 8 hours of multichannel PSG takes hours of tedious manual scoring and collapses complex neurobiology into 5 discrete stages. SleepLens solves this with a two-tiered AI system: an ensemble staging model with calibrated uncertainty, and a continuous sleep depth transformer with a 96-feature phenotyping engine, deployed in a production-ready clinical web application.")

    # =========================================================================
    # SLIDE 2: Dataset, Preprocessing & Standardization
    # =========================================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "Rigorous Benchmark: Sleep-EDF Expanded Standards", timing_text="0:45 – 1:30 | 45s")

    # Left Card: Dataset Stats
    add_card(s2, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    tb_ds = s2.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(4.9))
    tf_ds = tb_ds.text_frame
    tf_ds.word_wrap = True
    
    p = tf_ds.paragraphs[0]
    p.text = "Sleep-EDF Database Expanded (PhysioNet)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    
    p = tf_ds.add_paragraph()
    p.text = "• 197 All-Night PSG Recordings (~2,000 hours of continuous data)\n• 237,936 Standardized 30-Second Epochs (epochs_v3 benchmark)\n• Dual Cohort Architecture:\n   - Sleep Cassette (SC): 153 recordings (ambulatory healthy subjects)\n   - Sleep Telemetry (ST): 44 recordings (hospital sleep clinic cohort)\n• Gold Standard: Expert visual scoring (Wake, N1, N2, N3, REM)\n• Leakage-Free Validation: Subject-atomic 5-fold CV (folds5_v3)\n   Strict zero-subject-overlap across training, tuning, and evaluation."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(10)

    # Right Card: Preprocessing & Harmonization Pipeline
    add_card(s2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    tb_pre = s2.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.9))
    tf_pre = tb_pre.text_frame
    tf_pre.word_wrap = True
    
    p = tf_pre.paragraphs[0]
    p.text = "Signal Harmonization & Windowing Pipeline"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_pre.add_paragraph()
    p.text = "1. Canonical Resampling (100 Hz):\n   Bilinear/polyphase anti-aliasing resampling across all input channels.\n\n2. Channel Normalization & Mapping:\n   • Frontal EEG: EEG Fpz-Cz (microvolt standardized)\n   • Occipital/Parietal EEG: EEG Pz-Oz\n   • Ocular Dipole: EOG horizontal\n   • Chin Muscle Tone: EMG submental (sub-envelope & RMS)\n\n3. Lights-Off Windowing (BENCHMARK_30):\n   Trims massive awake pre/post recording padding to isolate true sleep opportunity period while preserving full transition dynamics."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(10)

    add_notes(s2, "Target: 0:45 - 1:30 (45s)\nTo ensure clinical validity, we used the full Sleep-EDF Expanded dataset across all 197 nights and nearly 238,000 epochs. Unlike many published papers that suffer from subject leakage, we enforced a strict subject-atomic 5-fold cross-validation scheme. We evaluated both cohorts: ambulatory cassette and clinical telemetry. Every recording is harmonized to 100 Hz, with standardized electrode alignments and automated lights-off windowing to ensure fair, reproducible benchmarking.")

    # =========================================================================
    # SLIDE 3: Tier 1 Model — Staging Benchmarks
    # =========================================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "Tier 1: 30-Second Epoch Sleep Stage Classification (SSC)", timing_text="1:30 – 2:25 | 55s")

    # Table of Results
    rows, cols = 7, 10
    tbl_shape = s3.shapes.add_table(rows, cols, Inches(0.8), Inches(1.5), Inches(11.733), Inches(3.2))
    table = tbl_shape.table

    col_widths = [Inches(1.0), Inches(2.7), Inches(1.0), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9), Inches(1.0)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    headers = ["ID", "Architecture / Sensors", "Macro-F1", "SC F1", "ST F1", "N1 F1", "N2 F1", "N3 F1", "REM F1", "Cohen's κ"]
    data = [
        ["E19", "4-Model Diverse Ensemble", "0.8425", "0.836", "0.865", "0.635", "0.892", "0.833", "0.915", "0.831"],
        ["E11", "AnySleep 3ch (Fpz+Pz+EOG)", "0.8267", "0.809", "0.878", "0.622", "0.880", "0.788", "0.904", "0.812"],
        ["E11d", "AnySleep 2-Sensor (Fpz+EOG)", "0.8227", "0.811", "0.865", "0.592", "0.880", "0.812", "0.905", "0.808"],
        ["E20", "RobustSleepNet (Clean - No EDF)", "0.7501", "0.736", "0.788", "0.442", "0.828", "0.717", "0.855", "0.724"],
        ["E00", "LightGBM Bandpower (11 Feats)", "0.6672", "—", "—", "0.339", "0.752", "0.687", "0.585", "0.612"],
        ["E01", "YASA Heuristic Baseline (Fpz)", "0.5225", "0.486", "0.611", "0.100", "0.737", "0.538", "0.557", "0.473"],
    ]

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 41, 59)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT

    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = RGBColor(19, 78, 74) # Highlight winning
            else:
                cell.fill.fore_color.rgb = RGBColor(15, 23, 42)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_WHITE if r == 0 else TEXT_MUTED
            if r == 0:
                p.font.bold = True

    # Analysis Card below table
    add_card(s3, Inches(0.8), Inches(4.9), Inches(11.733), Inches(2.0))
    tb_an = s3.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.8))
    tf_an = tb_an.text_frame
    tf_an.word_wrap = True
    
    p = tf_an.paragraphs[0]
    p.text = "Key Scientific Findings in Model Exploration:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = tf_an.add_paragraph()
    p.text = "1. Ensemble Diversity Beats Homogeneous Stacking: Adding our lightweight LightGBM model (0.667 F1) to the deep neural models boosted the ensemble by +0.016 F1 because it operates on orthogonal spectral hand-crafted features rather than raw waveforms.\n2. Record Class Scores: The winning E19 system sets study records across all 4 key sleep stages: N1 (0.635), N2 (0.892), N3 (0.833), and REM (0.915).\n3. Rejection of Viterbi Decoding: Applying post-hoc Viterbi transition decoding degraded scores (0.8267 -> 0.8259) because AnySleep already possesses an intrinsic 14-minute temporal receptive field."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(4)

    add_notes(s3, "Target: 1:30 - 2:25 (55s)\nFor sleep stage classification, we systematically benchmarked 25+ configurations. Classical heuristics like YASA struggle severely on N1 sleep at only 0.10 F1. Deep neural models like AnySleep perform exceptionally at 0.8267. Our winning system, E19, is a 4-member fold-honest ensemble combining AnySleep 3-channel, AnySleep 2-channel, U-Sleep CSDP, and our engineered LightGBM model. It reached 0.8425 Macro-F1 with a Cohen's Kappa of 0.831. Crucially, the weak LightGBM model improved the ensemble by introducing non-redundant spectral feature diversity.")

    # =========================================================================
    # SLIDE 4: Sensor Ablation & Interpretability
    # =========================================================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "SSC Insights: Wearable Sensor Ablation & Explainability", timing_text="2:25 – 3:15 | 50s")

    # Left: Channel Ablation Card
    add_card(s4, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    tb_ab = s4.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(4.9))
    tf_ab = tb_ab.text_frame
    tf_ab.word_wrap = True
    
    p = tf_ab.paragraphs[0]
    p.text = "Minimal Sensor Montage Discovery"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = tf_ab.add_paragraph()
    p.text = "Evaluating hardware reduction for home wearable integration:\n\n• Full 3-Channel (Fpz-Cz + Pz-Oz + EOG): 0.8267 F1 (Baseline)\n• 2-Sensor (Fpz-Cz + EOG): 0.8227 F1 (-0.004 Δ!)\n   Retains 99.5% of full system performance with only 2 leads!\n• Dual EEG (Fpz-Cz + Pz-Oz, no EOG): 0.8134 F1 (-0.013 Δ)\n• Single Frontal EEG (Fpz-Cz): 0.8015 F1 (-0.025 Δ)\n• Single Occipital EEG (Pz-Oz): 0.7736 F1 (-0.053 Δ)\n\nKey Takeaway: The Pz-Oz channel is completely droppable. Adding EOG to frontal EEG achieves higher true REM recall (0.913 vs 0.897) by directly capturing ocular dipole saccades."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(8)

    # Right: Interpretability & Confidence Card
    add_card(s4, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    tb_int = s4.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.9))
    tf_int = tb_int.text_frame
    tf_int.word_wrap = True
    
    p = tf_int.paragraphs[0]
    p.text = "Attention Physiology & Doctor-in-the-Loop"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_int.add_paragraph()
    p.text = "1. Physiological Channel Attention Routing:\n   Skip-connection attention weights match human clinical rules:\n   • Wake: 67.2% attention on Occipital EEG (Pz) (alpha rhythm)\n   • N3: 93.0% attention on EEG (Fpz+Pz), eyes quiet (<7.1% EOG)\n   • REM: Attention surges to 52.0% on EOG (saccadic eye bursts)\n\n2. Monotonic Confidence Calibration (ECE = 0.029):\n   • Conf ≥ 0.90 (30.3% of night): 99.5% empirical accuracy\n   • Conf 0.70–0.89: 90.4% accuracy\n   • Conf < 0.60: 53.6% accuracy -> Automated Review Flagging!\n   Only ~15% of borderline epochs require clinician review."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(8)

    add_notes(s4, "Target: 2:25 - 3:15 (50s)\nTwo practical questions: Can we run this on a minimal home device, and is the model interpretable? Our ablation proved that a two-sensor setup—one frontal EEG and one EOG—achieves 0.8227 F1, losing only 0.004 compared to three channels while actually improving REM detection. Deep skip-layer attention reveals the network dynamically shifts attention to posterior alpha in Wake, slow waves in N3, and surges to 52% EOG attention during REM. Finally, our confidence calibration is monotonic: epochs above 90% confidence are 99.5% accurate, enabling intelligent human-in-the-loop triage.")

    # =========================================================================
    # SLIDE 5: Tier 2 Model — Continuous Depth (SDI) & 96 Features
    # =========================================================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "Tier 2: Continuous Sleep Depth (SDI) & 96-Feature Phenotyping", timing_text="3:15 – 4:05 | 50s")

    # Top Card: Beyond Discrete Hypnograms
    add_card(s5, Inches(0.8), Inches(1.5), Inches(11.733), Inches(2.2))
    tb_sdi = s5.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.1), Inches(2.0))
    tf_sdi = tb_sdi.text_frame
    tf_sdi.word_wrap = True
    
    p = tf_sdi.paragraphs[0]
    p.text = "SDI Transformer (Zhou et al., npj Digital Medicine 2025)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = tf_sdi.add_paragraph()
    p.text = "• Moving beyond 5 flat discrete stages to continuous neurobiological depth:\n   Computes a continuous Sleep Depth Index (0.0 = Light/Wake to 1.0 = Deep SWA Sink) per 30-s epoch.\n• Captures gradual transitions, micro-arousals, and cyclic alternating patterns (CAP) invisible in discrete hypnograms.\n• Standardized Research Vector (Zhou et al.):\n   - RB (Shallow Burden: SDI < 0.20)  |  AP (Mean Sleep Depth)  |  CV (Depth Instability)\n   - MDR (Mean REM Depth)            |  PR (Predicted REM Share)"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(6)

    # Bottom Card: 96 Clinical Feature Groups
    add_card(s5, Inches(0.8), Inches(3.9), Inches(11.733), Inches(2.9))
    tb_feat = s5.shapes.add_textbox(Inches(1.1), Inches(4.0), Inches(11.1), Inches(2.7))
    tf_feat = tb_feat.text_frame
    tf_feat.word_wrap = True
    
    p = tf_feat.paragraphs[0]
    p.text = "High-Dimensional Sleep Quality Measurement Engine (~96 Parameters)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_feat.add_paragraph()
    p.text = "• Continuity & Sleep Efficiency: TST, TIB, SE%, SOL, WASO, REM Latency, Awakenings, Sleep Fragmentation Index (SFI).\n• Macro-Architecture & Dynamics: Stage % of TST, 1st vs 2nd half N1/REM shift, 5x5 Markov Stage Transition Matrix.\n• Spectral Power Dynamics: Absolute & relative Welch power (Delta, Theta, Alpha, Sigma, Beta), SWA, Spectral/Permutation Entropy.\n• Sleep Microstructure: YASA Spindle Density & amplitude (N2), Slow Wave Amplitude & duration (NREM).\n• EMG & Physiological Autonomic: Submental EMG RMS, REM atonia ratio, Movement index, Respiration apnea proxy."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(6)

    add_notes(s5, "Target: 3:15 - 4:05 (50s)\nIn Tier 2, we transcend discrete 5-stage hypnograms. We implement the Sleep Depth Index Transformer from Zhou et al., published in npj Digital Medicine 2025. This computes a continuous 0.0 to 1.0 sleep depth score for every epoch, capturing micro-arousals and depth dynamics that discrete staging obscures. In tandem, our feature extraction engine computes ~96 clinical parameters per night: sleep continuity, 5x5 Markov transition matrices, spectral bandpowers, and microstructural events like YASA sleep spindle density and REM muscle atonia ratios.")

    # =========================================================================
    # SLIDE 6: Sleep Quality Index & Multi-Modal Scoring
    # =========================================================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Sleep Quality Index: Unified Scoring & PSQI Synthesis", timing_text="4:05 – 4:55 | 50s")

    # Left: Score Formulation
    add_card(s6, Inches(0.8), Inches(1.5), Inches(6.5), Inches(5.3))
    tb_sc = s6.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.9), Inches(4.9))
    tf_sc = tb_sc.text_frame
    tf_sc.word_wrap = True
    
    p = tf_sc.paragraphs[0]
    p.text = "SleepLens Unified Sleep Score (0–100)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = tf_sc.add_paragraph()
    p.text = "A transparent, clinically grounded 100-point composite score:\n\n• Sleep Efficiency (25 pts): Normalizes sleep time vs time in bed.\n• Continuous Sleep Depth (25 pts): Mean SDI cortical slow-wave saturation.\n• Sleep Continuity (20 pts): Penalizes stage shift index & awakenings.\n• Deep Sleep Share (10 pts): Physical restoration (N3 % of TST).\n• REM Sleep Share (10 pts): Cognitive restoration (REM % of TST).\n• Sleep Onset Latency (5 pts): Time required to transition to sleep.\n• Fragmentation Burden (5 pts): Resistance to shallow sleep intrusions.\n\nGrading: 85–100 Excellent | 70–84 Good | 55–69 Fair | <55 Poor"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(8)

    # Right: PSQI & Clinical Synthesis
    add_card(s6, Inches(7.6), Inches(1.5), Inches(4.9), Inches(5.3))
    tb_psqi = s6.shapes.add_textbox(Inches(7.9), Inches(1.7), Inches(4.3), Inches(4.9))
    tf_psqi = tb_psqi.text_frame
    tf_psqi.word_wrap = True
    
    p = tf_psqi.paragraphs[0]
    p.text = "Subjective + Objective Integration"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_psqi.add_paragraph()
    p.text = "• Pittsburgh Sleep Quality Index (PSQI):\n   7 component questionnaire (0–21 global score) fully digitized.\n\n• Paradoxical Insomnia Detection:\n   Cross-referencing subjective PSQI with objective PSG features identifies Sleep State Misperception (normal PSG architecture despite poor subjective complaints).\n\n• LLM Clinical Report Synthesis:\n   Automated narrative generator (gpt-6-luna) translates 96 features, SDI depth, and PSQI into structured clinical findings."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(8)

    add_notes(s6, "Target: 4:05 - 4:55 (50s)\nTo synthesize these complex metrics into an actionable patient indicator, we designed the SleepLens Unified Sleep Score. Rather than an uninterpretable black-box score, it is a transparent 100-point scale: 25% sleep efficiency, 25% continuous depth, 20% continuity, and the remainder distributed across N3 deep sleep, REM, latency, and fragmentation. We also incorporate the standardized 7-component PSQI questionnaire. By contrasting subjective perception against objective PSG signals, SleepLens can flag paradoxical insomnia and feeds all metrics into an automated clinical reporting engine.")

    # =========================================================================
    # SLIDE 7: End-to-End System Architecture Schematic
    # =========================================================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "End-to-End Architecture: Ingestion to Clinical Decision", timing_text="4:55 – 5:45 | 50s")

    # Workflow schematic cards horizontally
    step_data = [
        ("01 • Ingestion & Auth", "• Drag & drop EDF (≤500 MB)\n• Optional 7-part PSQI form\n• Patient metadata vinculation\n• JWT httpOnly security", TEAL_PRIMARY, Inches(0.8)),
        ("02 • Signal Prep", "• Canonical 100 Hz resample\n• Lights-off windowing\n• Channel routing & QC\n• Celery + Redis queue", CYAN_ACCENT, Inches(3.8)),
        ("03 • Dual AI Pipeline", "• Tier 1: 4-Model Ensemble\n  (Hypnogram + Confidence)\n• Tier 2: SDI Transformer Net\n  (Depth trajectory 0–1)", BLUE_LIGHT, Inches(6.8)),
        ("04 • Clinical UI & AI", "• SleepLens Score (0–100)\n• 96-Feature breakdown\n• LLM Diagnostic Report\n• OpenCode Chat Assistant", EMERALD_GREEN, Inches(9.8)),
    ]

    for title, desc, color, left_pos in step_data:
        add_card(s7, left_pos, Inches(1.5), Inches(2.733), Inches(3.2))
        tb = s7.shapes.add_textbox(left_pos + Inches(0.15), Inches(1.7), Inches(2.433), Inches(2.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(8)

    # Technology Stack Bar
    add_card(s7, Inches(0.8), Inches(4.9), Inches(11.733), Inches(1.9))
    tb_st = s7.shapes.add_textbox(Inches(1.1), Inches(5.0), Inches(11.1), Inches(1.7))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True
    p = tf_st.paragraphs[0]
    p.text = "Production Technology Stack & Data Contracts"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p2 = tf_st.add_paragraph()
    p2.text = "• Web Frontend: Next.js 14 App Router, TypeScript, Tailwind CSS, Lucide icons, React Query, Zustand.\n• Backend Core: Django REST Framework, PostgreSQL, Celery asynchronous workers, Redis broker.\n• Machine Learning Core: PyTorch, MNE-Python, LightGBM, YASA, SciPy signal processing.\n• Intelligence Layer: OpenAI API (gpt-6-luna) for diagnostic narrative generation; OpenCode local agent for consultation."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(4)

    add_notes(s7, "Target: 4:55 - 5:45 (50s)\nHere is our production system architecture. A clinician uploads an EDF file and optional PSQI through the Next.js web application. The Django REST backend sanitizes the upload and delegates heavy computation to Celery workers backed by Redis. Signals are canonically harmonized to 100 Hz, passed in parallel through both AI tiers—the 4-model staging ensemble and the SDI transformer—and combined with the 96-feature engine. Results are indexed in PostgreSQL and rendered back to the clinician with an LLM clinical report and an OpenCode consultation assistant.")

    # =========================================================================
    # SLIDE 8: Clinical Web Application Walkthrough
    # =========================================================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "Clinical Web Application: SleepLens UI / UX", timing_text="5:45 – 6:30 | 45s")

    # 4 UI Feature Cards
    ui_cards = [
        ("ScoreCard & KPIs", "• SleepLens Score (0–100) gauge\n• 7-component radar breakdown\n• Core KPIs: TST, SE%, SOL, WASO, REM Latency, Mean Confidence", Inches(0.8), Inches(1.5)),
        ("Interactive Hypnogram", "• Synchronized 30-s stage timeline\n• Confidence bands (High/Med/Low)\n• One-click review flags for borderline epochs (<60% confidence)", Inches(6.8), Inches(1.5)),
        ("SDI Depth & Signals", "• Continuous 0.0–1.0 depth curve\n• Interactive epoch raw signal explorer\n• Multi-channel review (EEG, EOG, EMG) with zero client latency", Inches(0.8), Inches(3.9)),
        ("AI Reports & Consultation", "• Automated Markdown report with clinician edit & sign-off\n• Context-injected OpenCode chat for diagnostic questions and second opinions", Inches(6.8), Inches(3.9)),
    ]

    for title, desc, left, top in ui_cards:
        add_card(s8, left, top, Inches(5.7), Inches(2.1))
        tb = s8.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.3), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEAL_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(6)

    # Bottom Callout
    add_card(s8, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.9), bg_color=RGBColor(19, 78, 74), border_color=TEAL_PRIMARY)
    tb_c = s8.shapes.add_textbox(Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.7))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "Designed for Sleep Technologists & Clinicians: Zero-lag review, automated triage, and auditable decisions."
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    add_notes(s8, "Target: 5:45 - 6:30 (45s)\nThis slide showcases the clinician interface. The dashboard presents the SleepLens Score and core KPIs at a glance. Below, the interactive hypnogram is coupled with real-time confidence tracking, highlighting exactly which epochs need human verification. Clinicians can view continuous sleep depth, open the raw signal viewer to inspect EEG microvolts for any 30-second epoch, review the AI-drafted diagnostic report, or ask questions in the OpenCode consultation chat. It reduces scoring overhead from 90 minutes down to a quick verification.")

    # =========================================================================
    # SLIDE 9: Clinical Impact & Roadmap
    # =========================================================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9)
    add_header(s9, "Deployment Performance, Clinical Impact & Future Roadmap", timing_text="6:30 – 7:00 | 30s")

    # Left: Deployment Benchmarks
    add_card(s9, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    tb_bm = s9.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(4.9))
    tf_bm = tb_bm.text_frame
    tf_bm.word_wrap = True
    
    p = tf_bm.paragraphs[0]
    p.text = "CPU Deployment Benchmarks (8 Threads)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY

    p = tf_bm.add_paragraph()
    p.text = "No Expensive GPU Required for Deployment:\n\n• Single Epoch Inference: 3.7–4.7 ms (x6,400 faster than real time)\n• Full 22-Hour PSG Night (2,650 epochs):\n   - EDF load & resample: 0.6 seconds\n   - Ensemble staging: 12.5 seconds\n   - SDI depth inference: 8.2 seconds\n   - 96-feature extraction: 9.8 seconds\n   - Total turnaround: < 35 seconds per night!\n\nClinical Productivity Gain: Reduces manual physician scoring time by >85% while providing deeper physiological insight."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(8)

    # Right: Roadmap
    add_card(s9, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    tb_rd = s9.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.9))
    tf_rd = tb_rd.text_frame
    tf_rd.word_wrap = True
    
    p = tf_rd.paragraphs[0]
    p.text = "Clinical Impact & Future Roadmap"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_rd.add_paragraph()
    p.text = "1. Wearable Surrogate Integration:\n   Deploying our trained ResUNet and WatchSleepNet models for wrist-worn actigraphy and PPG without any EEG leads.\n\n2. Disorder-Specific Phenotyping:\n   Expanding the proxy engine to full AASM Apnea-Hypopnea Index (AHI) and Periodic Limb Movement Disorder (PLMD) scoring.\n\n3. Multi-Center Clinical Trials:\n   Evaluating generalization across clinical cohorts with sleep apnea, insomnia, and Parkinsonian REM sleep behavior disorder (RBD)."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(8)

    add_notes(s9, "Target: 6:30 - 7:00 (30s)\nIn summary, SleepLens delivers state-of-the-art staging accuracy of 0.8425 F1 and 0.831 Kappa, with a complete turnaround time under 35 seconds on standard CPU hardware. By uniting discrete stage classification, continuous depth modeling, and automated clinical reporting, we cut physician scoring time by over 85% while providing unprecedented diagnostic depth. Our minimal two-sensor finding also enables accurate home wearable monitoring. SleepLens bridges machine learning with the daily realities of clinical sleep medicine. Thank you.")

    output_path = "SleepLens_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_deck()
